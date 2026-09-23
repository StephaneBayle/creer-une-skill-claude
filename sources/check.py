"""Contrôles des livrables. Usage : .venv/bin/python check.py"""
import sys, zipfile, re, json, subprocess
from pathlib import Path
import cv2, numpy as np
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
PPTX = ROOT / "skills-claude-presentation.pptx"
URLS = {"https://www.linkedin.com/in/stephanebayle/", "https://github.com/StephaneBayle",
        "https://github.com/StephaneBayle/creer-une-skill-claude",
        "https://academy.claude.com/fr/use-cases/package-your-brand-guidelines-in-a-skill"}
TARGET_MIN, TARGET_MAX = 45, 58  # durée visée, en minutes
MAX_WORDS = 50  # mots à l'écran, pastille et pied de page compris
# maquettes et tableaux denses : le texte y fait partie de l'image
MOCKUP_KINDS = {"hook", "anatomy", "catalog", "demo", "codedemo", "proscons", "usecase", "planmode", "creatorloop", "creatormatrix", "coworkcode"}
KINDS = json.loads(subprocess.check_output(["node", "-e", "console.log(JSON.stringify(require('./content').SLIDES.map(s => s.kind)))"], cwd=Path(__file__).parent))
errors = []

prs = Presentation(PPTX)
total = 0
for i, slide in enumerate(prs.slides, 1):
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    if len(notes.split()) < 20:
        errors.append(f"diapo {i} : notes absentes ou trop courtes")
    m = re.search(r"TIMING — (\d+):(\d+)", notes)
    if m:
        total += int(m[1]) * 60 + int(m[2])
    words = sum(len(sh.text_frame.text.split()) for sh in slide.shapes if sh.has_text_frame)
    if words > MAX_WORDS and KINDS[i - 1] not in MOCKUP_KINDS:
        errors.append(f"diapo {i} : {words} mots à l'écran (> {MAX_WORDS})")

print(f"{len(prs.slides)} diapos, timing cumulé {total // 60}:{total % 60:02d}")
if len(KINDS) != len(prs.slides):
    errors.append("content.js et le pptx n'ont pas le même nombre de diapos")
if not TARGET_MIN * 60 <= total <= TARGET_MAX * 60:
    errors.append(f"timing hors de {TARGET_MIN}–{TARGET_MAX} min")

# QR codes de la dernière diapo
found = set()
with zipfile.ZipFile(PPTX) as z:
    for name in z.namelist():
        if name.startswith("ppt/media/") and name.endswith(".png"):
            img = cv2.imdecode(np.frombuffer(z.read(name), np.uint8), cv2.IMREAD_COLOR)
            if img is None or img.shape[0] < 400:
                continue
            data, _, _ = cv2.QRCodeDetector().detectAndDecode(img)
            if data:
                found.add(data)
print("QR décodés :", sorted(found))
if found != URLS:
    errors.append(f"QR codes attendus {URLS}, trouvés {found}")

# Coordonnées présentes dans le document distribué
z = zipfile.ZipFile(ROOT / "skills-claude-notes-presentateur.docx")
doc = z.read("word/document.xml").decode() + z.read("word/_rels/document.xml.rels").decode()
for needle in ["stephane.bayle@gmail.com", *URLS]:
    if needle not in doc:
        errors.append(f"docx : {needle} absent")

# Accessibilité du pptx
with zipfile.ZipFile(PPTX) as z:
    slide_xml = [z.read(n).decode() for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)]
if any('lang="en-US"' in x for x in slide_xml):
    errors.append("pptx : texte encore déclaré en anglais (en-US)")
if any('descr="preencoded.png"' in x for x in slide_xml):
    errors.append("pptx : images sans texte alternatif ni marque décorative")
untitled = sum('<p:ph type="title"/>' not in x for x in slide_xml)
if untitled:
    errors.append(f"pptx : {untitled} diapos sans espace réservé titre")

# Contrastes de la charte (paires texte / fond déclarées dans tokens.json)
TK = json.loads((Path(__file__).with_name("tokens.json")).read_text())
def _lum(h):
    c = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]
for fg, bg in TK["textPairs"]:
    la, lb = sorted([_lum(TK["color"][fg]), _lum(TK["color"][bg])], reverse=True)
    ratio = (la + 0.05) / (lb + 0.05)
    if ratio < 4.5:
        errors.append(f"contraste {fg} sur {bg} : {ratio:.2f}:1 < 4.5")

# Document distribué
if 'w:lang w:val="fr-FR"' not in doc and 'w:val="fr-FR"' not in z.read("word/styles.xml").decode():
    errors.append("docx : langue française non déclarée")
ids = re.findall(r'<wp:docPr id="(\d+)"', doc)
if len(ids) != len(set(ids)):
    errors.append("docx : identifiants d'images en double")
import pymupdf
pdf = pymupdf.open(ROOT / "skills-claude-notes-presentateur.pdf")
if "fr-FR" not in pdf.xref_object(pdf.pdf_catalog()):
    errors.append("pdf : langue non déclarée")

for f in ["skills-claude-notes-presentateur.pdf", "cartons-quiz-ABCD.pdf"]:
    if not (ROOT / f).exists():
        errors.append(f"{f} manquant")

print("\n".join(errors) if errors else "Tous les contrôles passent.")
sys.exit(1 if errors else 0)
